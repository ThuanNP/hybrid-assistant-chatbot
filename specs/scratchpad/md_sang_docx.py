"""Chuyển bản nháp Markdown (cú pháp rút gọn trong _quy-uoc-noi-bo.md) sang tệp Word.

Thể thức trình bày theo Nghị định 30/2020/NĐ-CP (Phụ lục I):
- khổ A4, lề trên 2 cm, dưới 2 cm, trái 3 cm, phải 1,5 cm;
- phông Times New Roman, màu đen, nội dung cỡ 14, căn đều hai lề, lùi đầu dòng 1 cm,
  cách đoạn 6 pt, dãn dòng đơn;
- số trang ở giữa lề trên, cỡ 13, không hiển thị ở trang thứ nhất.
Khối lệnh, khối PROMPT và sơ đồ ASCII dùng phông đơn cách để giữ căn cột.

Cách dùng: python md_sang_docx.py <tieu_de_tai_lieu> <tep_ra.docx> <tep1.md> [<tep2.md> ...]
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT_CHU = "Times New Roman"
FONT_MA = "Consolas"
DEN = RGBColor(0, 0, 0)
CO_CHU = 14          # nội dung văn bản: 13-14
CO_BANG = 12         # chữ trong bảng
CO_MA = 9.5          # khối lệnh, khối PROMPT
CO_SO_TRANG = 13
RONG_VUNG_VIET = 16.5  # 21 cm - lề trái 3 cm - lề phải 1,5 cm

KIEU_KHOI = {
    # nhãn: (nền ô, tiêu đề trong ô, font mã?)
    "prompt": ("F2F2F2", None, True),
    "danhgia": ("F2F2F2", None, True),
    "batbuoc": ("FFFFFF", "BẮT BUỘC", False),
    "meo": ("FFFFFF", "MẸO", False),
    "canhbao": ("FFFFFF", "CẢNH BÁO", False),
}
NEN_MA = "F2F2F2"


def to_nen(o, mau):
    tcPr = o._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), mau)
    tcPr.append(shd)


def le_o(o, tren=80, duoi=80, trai=120, phai=120):
    tcPr = o._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for ten, gt in (("top", tren), ("bottom", duoi), ("start", trai), ("end", phai)):
        el = OxmlElement(f"w:{ten}")
        el.set(qn("w:w"), str(gt))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def lap_lai_tieu_de(hang):
    trPr = hang._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def dat_font(run, ten, co=None, dam=None, nghieng=None):
    run.font.name = ten
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for thuoc in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(thuoc), ten)
    if co:
        run.font.size = Pt(co)
    if dam is not None:
        run.bold = dam
    if nghieng is not None:
        run.italic = nghieng
    run.font.color.rgb = DEN


MAU_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def them_chu(doan, van_ban, co=None, dam_mac_dinh=False, nghieng=None):
    for phan in MAU_INLINE.split(van_ban):
        if not phan:
            continue
        if phan.startswith("**") and phan.endswith("**"):
            dat_font(doan.add_run(phan[2:-2]), FONT_CHU, co, True, nghieng)
        elif phan.startswith("`") and phan.endswith("`"):
            dat_font(doan.add_run(phan[1:-1]), FONT_MA, (co or CO_CHU) - 2, dam_mac_dinh, nghieng)
        else:
            dat_font(doan.add_run(phan), FONT_CHU, co, dam_mac_dinh, nghieng)


def khoang(doan, truoc=0, sau=6, don=True):
    pf = doan.paragraph_format
    pf.space_before = Pt(truoc)
    pf.space_after = Pt(sau)
    if don:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


class BoChuyen:
    def __init__(self, tieu_de):
        self.doc = Document()
        self.ten_tai_lieu = tieu_de
        self.so_h1 = 0
        self.trong_phan_c = False
        self._co_noi_dung_truoc = False
        self._thiet_lap()

    # ---------- thiết lập chung ----------
    def _thiet_lap(self):
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(1.5)
        sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
        sec.different_first_page_header_footer = True
        self.doc.core_properties.title = self.ten_tai_lieu

        st = self.doc.styles["Normal"]
        st.font.name = FONT_CHU
        st.font.size = Pt(CO_CHU)
        st.font.color.rgb = DEN
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CHU)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        for cap in (1, 2, 3, 4):
            h = self.doc.styles[f"Heading {cap}"]
            h.font.name = FONT_CHU
            h.font.size = Pt(CO_CHU)
            h.font.bold = True
            h.font.italic = cap == 4
            h.font.color.rgb = DEN
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.append(rf)
            for thuoc in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(thuoc), FONT_CHU)
            for thuoc in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if rf.get(qn(thuoc)) is not None:
                    del rf.attrib[qn(thuoc)]
            pf = h.paragraph_format
            pf.space_before = Pt(12 if cap <= 2 else 6)
            pf.space_after = Pt(6)
            pf.keep_with_next = True
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if cap == 1 else WD_ALIGN_PARAGRAPH.JUSTIFY
        for ten in ("TOC 1", "TOC 2", "TOC 3"):
            try:
                t = self.doc.styles[ten]
            except KeyError:
                continue
            t.font.name = FONT_CHU
            t.font.size = Pt(13)

        # số trang ở giữa lề trên, không hiển thị ở trang thứ nhất
        hd = sec.header.paragraphs[0]
        hd.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._truong(hd, "PAGE")

    def _truong(self, doan, ma):
        r = doan.add_run()
        dat_font(r, FONT_CHU, CO_SO_TRANG, False)
        b = OxmlElement("w:fldChar")
        b.set(qn("w:fldCharType"), "begin")
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = ma
        s = OxmlElement("w:fldChar")
        s.set(qn("w:fldCharType"), "separate")
        v = OxmlElement("w:t")
        v.text = "1"
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), "end")
        for el in (b, t, s, v, e):
            r._r.append(el)

    def muc_luc(self):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dat_font(p.add_run("MỤC LỤC"), FONT_CHU, CO_CHU, True)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        b = OxmlElement("w:fldChar")
        b.set(qn("w:fldCharType"), "begin")
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = 'TOC \\o "1-3" \\h \\z \\u'
        s = OxmlElement("w:fldChar")
        s.set(qn("w:fldCharType"), "separate")
        v = OxmlElement("w:t")
        v.text = "Nhấn chuột phải vào đây và chọn Update Field (hoặc F9) để cập nhật mục lục."
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), "end")
        for el in (b, t, s, v, e):
            r._r.append(el)
        self.ngat_trang()

    def ngat_trang(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------- khối ----------
    def bia(self, dong):
        for _ in range(6):
            self.doc.add_paragraph()
        for i, d in enumerate(dong):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 0:
                co, dam = 14, True
            elif i in (1, 2):
                co, dam = 16, True
            else:
                co, dam = 14, False
            dat_font(p.add_run(d), FONT_CHU, co, dam)
            khoang(p, 0, 12 if i < 3 else 6)
        self.ngat_trang()
        self.muc_luc()

    def tieu_de(self, cap, van_ban):
        if cap == 1:
            self.so_h1 += 1
            self.trong_phan_c = van_ban.startswith("PHẦN C")
            if self.so_h1 > 1:
                self.ngat_trang()
            van_ban = van_ban.upper()
        elif cap == 2 and self.trong_phan_c and van_ban.startswith("Giai đoạn") and self._co_noi_dung_truoc:
            self.ngat_trang()
        h = self.doc.add_heading(level=cap)
        them_chu(h, van_ban, CO_CHU, True, cap == 4)
        self._co_noi_dung_truoc = True

    def doan_van(self, van_ban):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.0)
        them_chu(p, van_ban)
        khoang(p, 0, 6)

    def danh_sach(self, muc):
        # muc: list[(cap, loai, so, van_ban)]
        for cap, loai, so, vb in muc:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            # Nghị định 30: chữ đầu dòng lùi 1 cm, dòng sau về lề trái như lời văn; cấp con lùi thêm 0,75 cm
            pf.left_indent = Cm(0.75 * cap)
            pf.first_line_indent = Cm(1.0)
            if vb.startswith("☐"):
                them_chu(p, vb)
            else:
                dau = f"{so}. " if loai == "so" else ("- " if cap == 0 else "+ ")
                dat_font(p.add_run(dau), FONT_CHU, CO_CHU, False)
                them_chu(p, vb)
            khoang(p, 0, 6)

    def bang(self, hang):
        so_cot = max(len(h) for h in hang)
        tb = self.doc.add_table(rows=len(hang), cols=so_cot)
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(hang):
            for j in range(so_cot):
                o = tb.cell(i, j)
                le_o(o, 50, 50, 90, 90)
                vb = h[j] if j < len(h) else ""
                p = o.paragraphs[0]
                khoang(p, 0, 0)
                if i == 0:
                    to_nen(o, "D9D9D9")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    them_chu(p, vb, CO_BANG, True)
                else:
                    them_chu(p, vb, CO_BANG)
        lap_lai_tieu_de(tb.rows[0])
        trong_so = []
        for j in range(so_cot):
            do_dai = [len(h[j]) if j < len(h) else 0 for h in hang]
            trong_so.append(max(min(sum(do_dai) / len(do_dai), 90), 8) ** 0.8)
        tong = sum(trong_so)
        tb.autofit = False
        for j in range(so_cot):
            rong = Cm(RONG_VUNG_VIET * trong_so[j] / tong)
            for i in range(len(hang)):
                tb.cell(i, j).width = rong
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def khoi(self, nhan, dong):
        while dong and not dong[-1].strip():
            dong.pop()
        if nhan == "prompt":
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(1.0)
            dat_font(p.add_run("PROMPT"), FONT_CHU, CO_CHU, True)
            khoang(p, 0, 6)
            p.paragraph_format.keep_with_next = True
        nen, td, la_ma = KIEU_KHOI.get(nhan, (NEN_MA, None, True))
        tb = self.doc.add_table(rows=1, cols=1)
        tb.style = "Table Grid"
        o = tb.cell(0, 0)
        o.width = Cm(RONG_VUNG_VIET)
        to_nen(o, nen)
        le_o(o)
        p0 = o.paragraphs[0]
        dau = True
        if td:
            dat_font(p0.add_run(td), FONT_CHU, CO_CHU, True)
            khoang(p0, 0, 3)
            dau = False
        for d in dong:
            p = p0 if dau else o.add_paragraph()
            dau = False
            if la_ma:
                dat_font(p.add_run(d.replace("\t", "    ")), FONT_MA, CO_MA)
                khoang(p, 0, 0)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                them_chu(p, d, 13)
                khoang(p, 0, 3)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # ---------- phân tích ----------
    def xu_ly(self, noi_dung):
        dong = noi_dung.splitlines()
        i = 0
        n = len(dong)
        doan = []

        def xa_doan():
            if doan:
                self.doan_van(" ".join(x.strip() for x in doan))
                doan.clear()

        while i < n:
            d = dong[i]
            s = d.strip()
            m_rao = re.match(r"^```(\w*)\s*$", s)
            if m_rao:
                xa_doan()
                nhan = m_rao.group(1) or "text"
                i += 1
                khoi = []
                while i < n and not re.match(r"^```\s*$", dong[i].strip()):
                    khoi.append(dong[i])
                    i += 1
                i += 1
                if nhan == "bia":
                    self.bia([x.strip() for x in khoi if x.strip()])
                else:
                    self.khoi(nhan, khoi)
                continue
            m_h = re.match(r"^(#{1,4})\s+(.*)$", s)
            if m_h:
                xa_doan()
                self.tieu_de(len(m_h.group(1)), m_h.group(2).strip())
                i += 1
                continue
            if s.startswith("|"):
                xa_doan()
                hang = []
                while i < n and dong[i].strip().startswith("|"):
                    o = [c.strip() for c in dong[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{3,}:?", c) for c in o if c):
                        hang.append(o)
                    i += 1
                self.bang(hang)
                continue
            m_ds = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", d)
            if m_ds:
                xa_doan()
                muc = []
                while i < n:
                    m2 = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", dong[i])
                    if m2:
                        cap = min(len(m2.group(1)) // 2, 2)
                        loai = "so" if m2.group(2)[0].isdigit() else "cham"
                        so = m2.group(2)[:-1] if loai == "so" else ""
                        muc.append([cap, loai, so, m2.group(3).strip()])
                        i += 1
                    elif dong[i].strip() and dong[i].startswith("  ") and muc:
                        muc[-1][3] += " " + dong[i].strip()
                        i += 1
                    else:
                        break
                self.danh_sach(muc)
                continue
            if not s:
                xa_doan()
                i += 1
                continue
            doan.append(s)
            i += 1
        xa_doan()

    def luu(self, duong_dan):
        self.doc.save(duong_dan)


def main():
    tieu_de, ra, *vao = sys.argv[1:]
    bc = BoChuyen(tieu_de)
    for tep in vao:
        with open(tep, encoding="utf-8") as f:
            bc.xu_ly(f.read())
    bc.luu(ra)
    print("Đã ghi", ra)


if __name__ == "__main__":
    main()
